from random import randint
from docx import Document
from docx.shared import Pt
from docx.shared import Cm


def getrandomnumber():
    value = randint(0, 12)
    return value


def createmultiplication():
    val1 = getrandomnumber()
    val2 = getrandomnumber()
    textbuilt = str(val1) + " x " + str(val2) + " ="
    return textbuilt


def createdivision():
    val1 = getrandomnumber()
    val2 = getrandomnumber()
    while val1 == 0 or val2 == 0:
        val1 = getrandomnumber()
        val2 = getrandomnumber()
    res = val1 * val2
    textbuilt = str(res) + " ÷ " + str(val2) + " ="
    return textbuilt


docu = Document()
table = docu.add_table(rows=33, cols=3)

style = docu.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(8)

sections = docu.sections
for section in sections:
    section.top_margin = Cm(1)

for i in range(99):
    t = randint(0, 1)
    if t == 0:
        task = createmultiplication()
    else:
        task = createdivision()
    if i < 33:
        cell = table.cell(i, 0)
        cell.text = task
    if i > 33 and i <= 66:
        cell = table.cell(i - 34, 1)
        cell.text = task
    if i > 66:
        cell = table.cell(i - 67, 2)
        cell.text = task

docu.save('test.doc')
